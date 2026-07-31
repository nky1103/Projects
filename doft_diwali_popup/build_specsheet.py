#!/usr/bin/env python3
"""Doft kiosk material specification & costing sheet (BOQ)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

wb = openpyxl.Workbook()
NAVY="4A343A"; BAND="6E4A50"; GOLD="EADFCE"; INPUT="FFF2CC"; GREEN="E2EFDA"; GREY="F5F1EA"; LINE="CFC4B4"
thin=Side(style="thin",color=LINE); box=Border(left=thin,right=thin,top=thin,bottom=thin)
INR=u'₹#,##0'; PCT='0.0%'

def st(c, val, bold=False, size=10, align="left", fill=None, color="2E2A28", wrap=False, fmt=None, border=True, italic=False):
    c.value=val; c.font=Font(name="Arial",size=size,bold=bold,color=color,italic=italic)
    c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if fill: c.fill=PatternFill("solid",fgColor=fill)
    if border: c.border=box
    if fmt: c.number_format=fmt
    return c

ws=wb.active; ws.title="Kiosk BOQ & Costing"
ws.sheet_view.showGridLines=False
# title
ws.merge_cells("A1:H1"); t=ws["A1"]; t.value="Doft Candles — Diwali Pop-Up Kiosk  |  Material Specification & Costing (per kiosk)"
t.font=Font(name="Arial",size=14,bold=True,color="FFFFFF"); t.fill=PatternFill("solid",fgColor=NAVY)
t.alignment=Alignment(horizontal="left",vertical="center",indent=1); ws.row_dimensions[1].height=28
ws.merge_cells("A2:H2"); s=ws["A2"]; s.value="8 ft × 8 ft island kiosk  ·  Prepared by ProMarcom  ·  Rates are indicative placeholders — confirm against vendor quotes"
s.font=Font(name="Arial",size=9,italic=True,color="FFFFFF"); s.fill=PatternFill("solid",fgColor=BAND)
s.alignment=Alignment(horizontal="left",vertical="center",indent=1); ws.row_dimensions[2].height=18
# legend
st(ws["A3"],"Legend:",bold=True,border=False)
st(ws["B3"],"Edit gold cells (Qty & Rate).",fill=INPUT,border=True)
st(ws["D3"],"White = formula",border=True); st(ws["E3"],"Green = totals",fill=GREEN,border=True)
for col in "CFGH": st(ws[f"{col}3"],"",border=False)

# header
hdr=5
heads=[("A","#",5),("B","Element",22),("C","Specification / Material & Finish",40),("D","Unit",7),("E","Qty",7),("F","Rate (₹)",12),("G","Amount (₹)",14),("H","Remarks",22)]
for col,txt,wdt in heads:
    st(ws[f"{col}{hdr}"],txt,bold=True,size=10,align="center",fill=NAVY,color="FFFFFF")
    ws.column_dimensions[col].width=wdt

sections=[
 ("1. Structure & Framework",[
   ("MS / GI frame & base structure","Welded frame, powder-coated","Lot",1,18000,"Core structure"),
   ("MDF / ply carcass (18 mm)","Walls, counter body","sqft",120,190,""),
   ("Base platform / deck framing","Raised floor structure","sqft",64,140,""),
 ]),
 ("2. Body Cladding & Finish",[
   ("PU-lacquered panels (ivory)","Smooth matt finish","sqft",140,260,"Body finish"),
   ("Fluted panel detailing","Counter / side fronts","sqft",30,320,""),
   ("Brass edge profiles & reveals","Brushed finish","rft",60,180,""),
 ]),
 ("3. Counter",[
   ("Engineered-marble top (20 mm)","Polished, ivory","sqft",18,650,"Billing counter"),
   ("Counter body & storage","Lockable shutters","Lot",1,22000,""),
   ("Under-counter drawers","Soft-close","Nos",3,3500,""),
 ]),
 ("4. Display & Shelving",[
   ("Backlit shelves / alcoves","Acrylic + LED, brass front","sqft",45,420,"Merchandising"),
   ("Toughened glass / acrylic shelves","8 mm","Nos",9,900,""),
   ("Display risers & props","Assorted","Lot",1,6000,""),
 ]),
 ("5. Branding & Signage",[
   ("Backlit fascia sign (DOFT)","Acrylic face, LED, brushed-gold letters","Nos",1,16000,"Hero signage"),
   ("Counter wordmark","Metal / acrylic letters","Nos",1,4500,""),
   ("Fascia framing & mounting","","Lot",1,3000,""),
 ]),
 ("6. Lighting",[
   ("Warm LED strips + drivers","3000K","rft",40,220,"Shelf & fascia"),
   ("Spotlights / track","Warm","Nos",6,900,""),
   ("Pendant lights (Option C)","Decorative","Nos",2,2500,"Optional"),
 ]),
 ("7. Electrical",[
   ("Concealed wiring & conduits","Copper","Lot",1,9000,"Single power drop"),
   ("Distribution board & MCBs","","Nos",1,4500,""),
   ("Sockets for billing / hardware","","Nos",4,350,""),
 ]),
 ("8. Flooring & Finishes",[
   ("Vinyl / laminate deck finish","Wood-look","sqft",64,120,""),
   ("Edge trims & skirting","Brass / anodised","rft",32,90,""),
 ]),
 ("9. Graphics & Styling",[
   ("Vinyl graphics & branding","Printed, laminated","sqft",25,140,""),
   ("Dried-flower & styling props","Diwali dressing","Lot",1,5000,""),
 ]),
 ("10. Hardware & Misc",[
   ("Locks, hinges, handles","Branded","Lot",1,4000,""),
   ("Fasteners & consumables","","Lot",1,2500,""),
   ("Fire extinguisher & safety","Compliance","Nos",1,2000,""),
 ]),
 ("11. Transport, Install & Dismantle",[
   ("Fabrication labour","","Lot",1,25000,""),
   ("Transport to mall (per city avg)","","Lot",1,12000,""),
   ("On-site install & finishing","","Lot",1,9000,""),
   ("Dismantle & return","End of season","Lot",1,6000,""),
 ]),
]

r=hdr+1; item=0; subtotal_rows=[]
for sec_name, rows in sections:
    ws.merge_cells(f"A{r}:H{r}")
    st(ws[f"A{r}"],sec_name,bold=True,size=10,fill=BAND,color="FFFFFF")
    r+=1
    first=r
    for (el,spec,unit,qty,rate,rem) in rows:
        item+=1
        zebra=GREY if item%2 else "FFFFFF"
        st(ws[f"A{r}"],item,align="center",fill=zebra,size=9)
        st(ws[f"B{r}"],el,fill=zebra,size=9,wrap=True)
        st(ws[f"C{r}"],spec,fill=zebra,size=9,wrap=True)
        st(ws[f"D{r}"],unit,align="center",fill=zebra,size=9)
        st(ws[f"E{r}"],qty,align="center",fill=INPUT,size=9)
        st(ws[f"F{r}"],rate,align="right",fill=INPUT,size=9,fmt=INR)
        st(ws[f"G{r}"],f"=E{r}*F{r}",align="right",fill=zebra,size=9,fmt=INR)
        st(ws[f"H{r}"],rem,fill=zebra,size=8,italic=True,wrap=True)
        r+=1
    last=r-1
    # subtotal
    st(ws[f"A{r}"],"",fill=GOLD); st(ws[f"B{r}"],f"Subtotal — {sec_name.split('. ',1)[1]}",bold=True,fill=GOLD,size=9)
    for col in "CDEF": st(ws[f"{col}{r}"],"",fill=GOLD)
    st(ws[f"G{r}"],f"=SUM(G{first}:G{last})",bold=True,align="right",fill=GOLD,size=9,fmt=INR)
    st(ws[f"H{r}"],"",fill=GOLD)
    subtotal_rows.append(r)
    r+=2

# ---- cost roll-up ----
def money_row(label, formula_or_val, fmt=INR, bold=False, fill=None):
    global r
    st(ws[f"A{r}"],"",fill=fill,border=True)
    st(ws[f"B{r}"],label,bold=bold,fill=fill,size=10)
    ws.merge_cells(f"B{r}:F{r}")
    st(ws[f"G{r}"],formula_or_val,bold=bold,align="right",fill=fill,size=10,fmt=fmt)
    st(ws[f"H{r}"],"",fill=fill,border=True)
    row=r; r+=1; return row

st(ws[f"B{r}"],"COST SUMMARY (per kiosk)",bold=True,size=11,fill=NAVY,color="FFFFFF"); ws.merge_cells(f"B{r}:H{r}")
for col in "A": st(ws[f"{col}{r}"],"",fill=NAVY)
r+=1
fab=money_row("Fabrication subtotal (all sections)","="+"+".join(f"G{sr}" for sr in subtotal_rows),bold=True,fill=GREEN)
# contingency %
st(ws[f"B{r}"],"Contingency",size=10); ws.merge_cells(f"B{r}:D{r}")
st(ws[f"A{r}"],"",border=True); st(ws[f"E{r}"],"",border=True)
st(ws[f"F{r}"],0.08,align="center",fill=INPUT,fmt=PCT); st(ws[f"G{r}"],f"=G{fab}*F{r}",align="right",fmt=INR); st(ws[f"H{r}"],"",border=True)
cont=r; r+=1
st(ws[f"B{r}"],"Design & project management",size=10); ws.merge_cells(f"B{r}:D{r}")
st(ws[f"A{r}"],"",border=True); st(ws[f"E{r}"],"",border=True)
st(ws[f"F{r}"],0.12,align="center",fill=INPUT,fmt=PCT); st(ws[f"G{r}"],f"=G{fab}*F{r}",align="right",fmt=INR); st(ws[f"H{r}"],"",border=True)
pm=r; r+=1
pre=money_row("Sub-total (pre-GST)",f"=G{fab}+G{cont}+G{pm}",bold=True,fill=GREEN)
st(ws[f"B{r}"],"GST",size=10); ws.merge_cells(f"B{r}:D{r}")
st(ws[f"A{r}"],"",border=True); st(ws[f"E{r}"],"",border=True)
st(ws[f"F{r}"],0.18,align="center",fill=INPUT,fmt=PCT); st(ws[f"G{r}"],f"=G{pre}*F{r}",align="right",fmt=INR); st(ws[f"H{r}"],"",border=True)
gst=r; r+=1
perk=money_row("PER-KIOSK TOTAL (incl. GST)",f"=G{pre}+G{gst}",bold=True,fill="D9B86B")
r+=1
# rollup
st(ws[f"B{r}"],"Number of kiosks",size=10); ws.merge_cells(f"B{r}:E{r}")
st(ws[f"A{r}"],"",border=True); st(ws[f"F{r}"],"",border=True)
st(ws[f"G{r}"],9,align="right",fill=INPUT,size=10); st(ws[f"H{r}"],"Edit to your rollout",italic=True,size=8)
nk=r; r+=1
prog=money_row("PROGRAMME TOTAL (all kiosks, incl. GST)",f"=G{perk}*G{nk}",bold=True,fill=NAVY,fmt=INR)
ws[f"G{prog}"].font=Font(name="Arial",size=11,bold=True,color="FFFFFF")
ws[f"B{prog}"].font=Font(name="Arial",size=11,bold=True,color="FFFFFF")
for col in "ABCDEFH": ws[f"{col}{prog}"].fill=PatternFill("solid",fgColor=NAVY)
ws[f"G{prog}"].fill=PatternFill("solid",fgColor=NAVY)

r+=2
st(ws[f"A{r}"],"Note: quantities and rates are indicative placeholders for planning. Confirm against production drawings and vendor quotes. Pendant lights apply to Option C only.",italic=True,size=8,border=False)
ws.merge_cells(f"A{r}:H{r}")

ws.freeze_panes="A6"
out="/home/user/Projects/doft_diwali_popup/Doft_Kiosk_Material_Spec_Costing.xlsx"
wb.save(out); print("Saved:",out)
